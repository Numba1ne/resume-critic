"""
File format conversion utilities (PDF ↔ DOCX).
"""
import os
import tempfile
from typing import Optional
from docx import Document
from docx.shared import Pt
from fpdf import FPDF


def pdf_to_docx_text(pdf_path: str, output_path: Optional[str] = None) -> str:
    """
    Convert PDF to DOCX by extracting text and creating DOCX.
    Note: This is a basic conversion - formatting may be lost.
    
    Args:
        pdf_path: Path to PDF file
        output_path: Optional path to save DOCX (if None, returns text only)
        
    Returns:
        Extracted text from PDF
    """
    from src.utils.pdf_handler import extract_pdf_text
    
    # Extract text from PDF
    with open(pdf_path, 'rb') as f:
        text = extract_pdf_text(f)
    
    if output_path:
        # Create DOCX from text
        doc = Document()
        for line in text.split('\n'):
            if line.strip():
                doc.add_paragraph(line.strip())
        doc.save(output_path)
    
    return text


def docx_to_pdf_basic(docx_path: str, output_path: str):
    """
    Convert DOCX to PDF using basic text extraction.
    Note: This is a simple conversion - use proper tools for production.
    
    Args:
        docx_path: Path to DOCX file
        output_path: Path to save PDF
    """
    from src.utils.docx_handler import read_docx, extract_docx_text
    
    # Read DOCX
    doc = read_docx(docx_path)
    text = extract_docx_text(doc)
    
    # Create PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', size=12)
    
    # Split text into lines and add to PDF
    for line in text.split('\n'):
        if line.strip():
            pdf.cell(0, 10, line.strip(), ln=1)
    
    pdf.output(output_path)


def extract_text_from_file(file_path: str) -> str:
    """
    Extract text from file (PDF or DOCX).
    
    Args:
        file_path: Path to file
        
    Returns:
        Extracted text
    """
    _, ext = os.path.splitext(file_path.lower())
    
    if ext == '.pdf':
        from src.utils.pdf_handler import extract_pdf_text
        with open(file_path, 'rb') as f:
            return extract_pdf_text(f)
    elif ext in ['.docx', '.doc']:
        from src.utils.docx_handler import read_docx, extract_docx_text
        doc = read_docx(file_path)
        return extract_docx_text(doc)
    else:
        raise ValueError(f"Unsupported file format: {ext}")
